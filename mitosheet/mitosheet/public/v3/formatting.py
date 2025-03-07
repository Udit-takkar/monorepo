from typing import Optional

from openpyxl.styles import Font, PatternFill
from openpyxl.styles import NamedStyle
from pandas import ExcelWriter

def add_formatting_to_excel_sheet(
        writer: ExcelWriter,
        sheet_name: str,
        header_background_color: Optional[str]=None,
        header_font_color: Optional[str]=None,
        even_background_color: Optional[str]=None,
        even_font_color: Optional[str]=None,
        odd_background_color: Optional[str]=None,
        odd_font_color: Optional[str]=None
    ) -> None:
        """
        Adds formatting to the sheet_name, based on the formatting the user
        currently has applied in the frontend. 

        NOTE: this is a Mito Pro feature.
        """
        workbook = writer.book
        sheet = workbook.get_sheet_by_name(sheet_name)
        
        # Add formatting to the header row   
        header_name = f"{sheet_name}_Header"
        header_format = NamedStyle(name=header_name)
        if header_font_color:
            # Remove the # from the color
            header_format.font = Font(color=header_font_color[1:])
        if header_background_color:
            # Remove the # from the color
            header_format.fill = PatternFill(start_color=header_background_color[1:], end_color=header_background_color[1:], fill_type="solid")
        
        # Only add format if there is a header color or background color
        has_header_formatting = header_background_color or header_font_color 
        if has_header_formatting:
            # Add named styles for the header rows to improve performance
            workbook.add_named_style(header_format)

            # Write the formatting to the sheet
            for col in range(1, sheet.max_column + 1):
                sheet.cell(row=1, column=col).style = header_name

        # Add formatting to the rows
        even_name = f"{sheet_name}_Even"
        odd_name = f"{sheet_name}_Odd"
        even_format = NamedStyle(name=even_name)
        odd_format = NamedStyle(name=odd_name)

        # Remove the # from the colors and define the formatting objects
        if even_background_color:
            even_format.fill = PatternFill(start_color=even_background_color[1:], end_color=even_background_color[1:], fill_type="solid")
        if even_font_color:
            even_format.font = Font(color=even_font_color[1:])
        if odd_background_color:
            odd_format.fill = PatternFill(start_color=odd_background_color[1:], end_color=odd_background_color[1:], fill_type="solid")
        if odd_font_color:
            odd_format.font = Font(color=odd_font_color[1:])
        
        # Only add format if there is a background color or font color
        has_row_formatting = even_background_color or even_font_color or odd_background_color or odd_font_color
        if has_row_formatting:
            workbook.add_named_style(even_format)
            workbook.add_named_style(odd_format)

            for row in range(2, sheet.max_row + 1):
                for col in range(1, sheet.max_column + 1):
                    if row % 2 == 0:
                        sheet.cell(row=row, column=col).style = even_name
                    else:
                        sheet.cell(row=row, column=col).style = odd_name